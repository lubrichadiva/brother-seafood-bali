from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import content as C

GOLD = RGBColor(0xC9, 0xA2, 0x27)
MAROON = RGBColor(0x8B, 0x1A, 0x1A)
DARK = RGBColor(0x1A, 0x15, 0x13)
MUTED = RGBColor(0x6E, 0x65, 0x5F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _run(p, text, size=10.5, bold=False, italic=False, color=DARK, font="Lato"):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return r


def _para(doc, text="", size=10.5, bold=False, italic=False, color=DARK, align=None, font="Lato", space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        _run(p, text, size, bold, italic, color, font)
    return p


def _gold_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p_pr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C9A227")
    bdr.append(bottom)
    p_pr.append(bdr)


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    _run(p, text, 14, bold=True, color=MAROON, font="Playfair Display")
    _gold_rule(doc)


def _bullets(doc, items, marker="✓"):
    for it in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.5)
        _run(p, f"{marker}  ", 10.5, bold=True, color=GOLD)
        _run(p, it, 10.5)


def _notes(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.5)
        _run(p, ln, 9, italic=True, color=MUTED)


def _table(doc, headers, rows, widths=None, price_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        _shade(cell, "8B1A1A")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(cell.paragraphs[0], h, 10, bold=True, color=WHITE)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            if ri % 2 == 1:
                _shade(cells[i], "FDFBF5")
            p = cells[i].paragraphs[0]
            is_price = price_col is not None and i == price_col
            if i > 0 or len(headers) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(p, str(val), 10, bold=is_price or i == 0 and len(headers) > 2, color=MAROON if is_price else DARK)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(proposal: dict, logo_bytes=None, gallery=None) -> BytesIO:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.8)
        s.left_margin = s.right_margin = Cm(2)
    doc.styles["Normal"].font.name = "Lato"
    doc.styles["Normal"].font.size = Pt(10.5)

    if logo_bytes:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(BytesIO(logo_bytes), width=Cm(5))

    _para(doc, C.BRAND["name"], 26, bold=True, color=MAROON, align="center", font="Playfair Display", space_after=2)
    _para(doc, C.BRAND["tagline_en"], 11, italic=True, color=GOLD, align="center", space_after=0)
    _para(doc, C.BRAND["tagline_zh"], 11, italic=True, color=GOLD, align="center", space_after=8)
    _gold_rule(doc)
    _para(doc, C.BRAND["doc_title"], 16, bold=True, color=DARK, align="center", font="Playfair Display", space_after=0)
    _para(doc, C.BRAND["doc_subtitle"], 11, bold=True, color=MUTED, align="center", space_after=10)

    contact = f" ({proposal.get('contact_type', 'WeChat')}: {proposal['client_contact']})" if proposal.get("client_contact") else ""
    for label, val in [
        ("Prepared for: ", f"{proposal['client_name']}{contact}"),
        ("Prepared by: ", proposal.get("prepared_by") or "Hendri (Ko Aby) — Owner, Brother Seafood Bali"),
        ("Date: ", proposal.get("date") or "_______________"),
    ]:
        p = _para(doc, align="center", space_after=1)
        _run(p, label, 10.5, bold=True, color=MAROON)
        _run(p, val, 10.5)

    _heading(doc, C.SECTIONS["about"])
    for tag, key in [("EN", "en"), ("中文", "zh"), ("ID", "id")]:
        p = _para(doc, space_after=6)
        _run(p, f"{tag}:  ", 10.5, bold=True, color=GOLD)
        _run(p, C.ABOUT[key], 10.5)

    _heading(doc, C.SECTIONS["location"])
    for ln in C.LOCATION:
        _para(doc, ln, 10.5, space_after=2)

    _heading(doc, C.SECTIONS["packages"])
    _para(doc, "Minimum 5 Pax · 最少 5 人 · Minimum 5 Pax", 10, italic=True, color=MUTED)
    _table(
        doc,
        ["Package 套餐", "Fish 鱼", "Calamari 鱿鱼", "Prawn 虾", "Clam 蛤蜊", "Price / Pax"],
        [[p["name"], p["fish"], p["calamari"], p["prawn"], p["clam"], f"IDR {p['price']:,}"] for p in C.PACKAGES],
        widths=[2.4, 2.4, 2.8, 2.4, 2.4, 3.6], price_col=5,
    )
    _notes(doc, C.PACKAGE_NOTES)
    _para(doc, "ALL PACKAGES INCLUDE · 所有套餐包含 · SEMUA PAKET SUDAH TERMASUK:", 11, bold=True, color=GOLD, font="Playfair Display")
    _bullets(doc, C.INCLUDES)

    _heading(doc, C.SECTIONS["menu"])
    _table(doc, ["Menu / Paket", "Price", "Pax", "Detail"], [[a,b,c,d] for a,b,c,d in C.MENU_HIGHLIGHTS], widths=[4.2,3.0,2.4,7.0], price_col=1)
    _para(doc, "Menu prices and availability are subject to the current menu and confirmation at booking.", 9, italic=True, color=MUTED)

    _heading(doc, C.SECTIONS["events"])
    _bullets(doc, C.EVENTS, marker="•")

    _heading(doc, C.SECTIONS["booking"])
    _para(doc, "Wedding, birthday, gathering, corporate and private events can be arranged with a custom quotation based on guest count and requested services.", 10.5)
    _table(doc, ["Event", "Booking arrangement"], [[a,b] for a,b in C.EVENT_BOOKING_DETAILS], widths=[5.2,11.4])
    _para(doc, "Booking Information Required · 预订资料 · DATA BOOKING", 11, bold=True, color=GOLD, font="Playfair Display")
    _bullets(doc, C.BOOKING_CHECKLIST, marker="•")

    _heading(doc, C.SECTIONS["commission"])
    _para(doc, "Commission Structure · 佣金结构 · Struktur Komisi", 11, bold=True, color=GOLD, font="Playfair Display")
    _table(doc, ["Item · 项目 · Item", "Commission · 佣金 · Komisi"], [list(r) for r in C.COMMISSION], widths=[12.5, 3.5], price_col=1)
    _notes(doc, C.COMMISSION_NOTES)
    _para(doc, "Complimentary for Tour Guide & Driver · 导游及司机免费待遇 · Kompliment untuk Guide & Driver", 11, bold=True, color=GOLD, font="Playfair Display")
    _bullets(doc, C.COMPLIMENTARY, marker="•")

    _heading(doc, C.SECTIONS["facilities"])
    _bullets(doc, C.FACILITIES, marker="•")
    _notes(doc, C.FACILITIES_NOTES)

    _heading(doc, C.SECTIONS["terms"])
    for i, t in enumerate(C.TERMS, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.5)
        _run(p, f"{i}.  ", 10.5, bold=True, color=MAROON)
        _run(p, t, 10.5)

    if proposal.get("notes"):
        _para(doc, "Special Notes · 备注 · Catatan Khusus", 11, bold=True, color=GOLD, font="Playfair Display")
        _para(doc, proposal["notes"], 10.5, italic=True)

    _heading(doc, C.SECTIONS["contact"])
    _para(doc, f"{C.BRAND['owner']} — Owner", 13, bold=True, color=MAROON, align="center", font="Playfair Display")
    for label, val in [("WhatsApp: ", C.BRAND["whatsapp"]), ("WeChat: ", C.BRAND["wechat"]), ("Instagram: ", C.BRAND["instagram"]), ("Address: ", C.BRAND["address"])]:
        p = _para(doc, align="center", space_after=1)
        _run(p, label, 10.5, bold=True, color=GOLD)
        _run(p, val, 10.5)

    if gallery:
        _heading(doc, C.SECTIONS["gallery"])
        t = doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i in range(0, len(gallery), 2):
            cells = t.add_row().cells
            for j, item in enumerate(gallery[i:i + 2]):
                img, cap = item
                p = cells[j].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    p.add_run().add_picture(BytesIO(img), width=Cm(7.6))
                except Exception:
                    continue
                if cap:
                    cp = cells[j].add_paragraph()
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _run(cp, cap, 9, italic=True, color=MUTED)

    doc.add_paragraph()
    _gold_rule(doc)
    _para(doc, C.CLOSING[0], 18, bold=True, color=MAROON, align="center", font="Playfair Display", space_after=4)
    _para(doc, C.CLOSING[1], 10.5, italic=True, color=MUTED, align="center", space_after=1)
    _para(doc, C.CLOSING[2], 10.5, italic=True, color=MUTED, align="center", space_after=16)

    _para(doc, "Signature & Stamp · 签名与盖章 · Tanda Tangan & Stempel", 10, bold=True, color=GOLD)
    doc.add_paragraph()
    doc.add_paragraph()
    _para(doc, "______________________________", 10.5, space_after=0)
    _para(doc, "Hendri (Ko Aby) — Owner, Brother Seafood Bali", 10, color=MUTED)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
