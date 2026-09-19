from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MAROON = "8B1A1A"
GOLD = "C9A227"
CREAM = "FDFBF5"
THIN = Side(style="thin", color="E5DECF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
IDR = '"IDR" #,##0'


def _header(ws, row, headers):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", name="Lato")
        c.fill = PatternFill("solid", fgColor=MAROON)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def _title(ws, text, month_label):
    ws["A1"] = "BROTHER SEAFOOD BALI"
    ws["A1"].font = Font(bold=True, size=16, color=MAROON, name="Playfair Display")
    ws["A2"] = text
    ws["A2"].font = Font(bold=True, size=12, color=GOLD)
    ws["A3"] = f"Period: {month_label}"
    ws["A3"].font = Font(italic=True, color="6E655F")


def _autowidth(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_report_xlsx(month_label: str, agents: list, visits: list, totals: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary per TA"
    _title(ws, "Monthly Commission Report · 月度佣金报告 · Laporan Komisi Bulanan", month_label)
    headers = ["Travel Agent", "Visits", "Pax", "Food Bill", "Beverage Bill", "Total Bill", "Commission", "Paid", "Unpaid"]
    _header(ws, 5, headers)
    r = 6
    for a in agents:
        vals = [a["ta_name"], a["visits"], a["pax"], a["food_bill"], a["beverage_bill"], a["bill"], a["commission"], a["paid"], a["unpaid"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            if i >= 4:
                c.number_format = IDR
            if (r - 6) % 2 == 1:
                c.fill = PatternFill("solid", fgColor=CREAM)
        ws.cell(row=r, column=1).font = Font(bold=True)
        r += 1
    tvals = ["TOTAL", totals["visits"], totals["pax"], totals["food_bill"], totals["beverage_bill"], totals["bill"], totals["commission"], totals["paid"], totals["unpaid"]]
    for i, v in enumerate(tvals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = Font(bold=True, color=MAROON)
        c.fill = PatternFill("solid", fgColor="F3E7B5")
        c.border = BORDER
        if i >= 4:
            c.number_format = IDR
    _autowidth(ws, [28, 8, 8, 16, 16, 16, 16, 16, 16])
    ws.freeze_panes = "A6"

    wd = wb.create_sheet("Visit Detail")
    _title(wd, "Visit Detail · 到访明细 · Detail Kunjungan", month_label)
    dh = ["Date", "Travel Agent", "Type", "Package", "Pax", "Food Bill", "Beverage Bill", "Uang Hadir", "Commission Food", "Commission Bev", "Total Commission", "Paid", "Notes"]
    _header(wd, 5, dh)
    r = 6
    for v in sorted(visits, key=lambda x: (x["date"], x["ta_name"])):
        rate = "Contract Rate" if v.get("rate_type", "contract") == "contract" else "À la carte"
        vals = [v["date"], v["ta_name"], rate, v["package"] if rate == "Contract Rate" else "-", v["pax"], v["food_bill"], v["beverage_bill"], v["uang_hadir"],
                v["commission_food"], v["commission_bev"], v["commission_total"], "Yes" if v.get("paid") else "No", v.get("notes", "")]
        for i, val in enumerate(vals, 1):
            c = wd.cell(row=r, column=i, value=val)
            c.border = BORDER
            if 6 <= i <= 11:
                c.number_format = IDR
        r += 1
    _autowidth(wd, [12, 24, 14, 9, 6, 15, 15, 13, 15, 15, 16, 7, 30])
    wd.freeze_panes = "A6"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
